"""Validate relational data, author-supplied aggregates, and saved RoB summaries."""
import csv
import math
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook
from export_inputs import ROOT, SHEETS

checks = []
def check(name, condition):
    checks.append((name, "PASS" if condition else "FAIL"))

def records(name):
    with (ROOT / "meta_analysis/data" / f"{name}.csv").open(encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))

studies, reports, conditions, outcomes = map(records, ["studies", "reports", "conditions", "outcomes"])
check("17 studies; 18 populations; 53 conditions; 370 outcomes",
      (len({r['study_id'] for r in studies}), len(studies), len(conditions), len(outcomes)) == (17,18,53,370))
for data, key in [(studies,'population_id'),(reports,'report_id'),(conditions,'condition_id'),(outcomes,'outcome_id')]:
    check(f"Unique {key}", len({r[key] for r in data}) == len(data) and all(r[key] for r in data))
study_ids = {r['study_id'] for r in studies}
populations = {r['population_id']:r['study_id'] for r in studies}
condition_map = {r['condition_id']:r for r in conditions}
check("Publication study links", all(r['study_id'] in study_ids for r in reports))
check("Condition and outcome population links", all(populations.get(r['population_id']) == r['study_id'] for r in conditions+outcomes))
check("Outcome condition links", all(not r['condition_id'] or
      (r['condition_id'] in condition_map and condition_map[r['condition_id']]['population_id'] == r['population_id']) for r in outcomes))
check("Derivation outcome links", all(not r['outcome_id'] or r['outcome_id'] in {o['outcome_id'] for o in outcomes} for r in records('derivations')))
workbook=load_workbook(ROOT/'extraction.xlsx',read_only=True,data_only=False)
for sheet, filename in SHEETS.items():
    with (ROOT/'meta_analysis/data'/f'{filename}.csv').open(encoding='utf-8',newline='') as f:
        actual=list(csv.reader(f))
    expected=[[str(v) if v is not None else '' for v in row] for row in workbook[sheet].values]
    check(f"Workbook/CSV agreement: {sheet}",actual==expected)
workbook.close()
author=ROOT/'meta_analysis/data/author_provided/Enomoto_2026/ttest_iAUC_1h5h.csv'
with author.open(encoding='utf-8-sig',newline='') as f:
    rows=[r for r in csv.reader(f) if r and r[0].startswith('day')]
check("Author data are four aggregate paired-test records",len(rows)==4 and all(len(r)==11 and r[1]=='13' and r[2]=='iAUC_1h - iAUC_5h' for r in rows))
for row in rows:
    day=row[0][-1]
    extracted=[o for o in outcomes if o['study_id']=='Enomoto_2026' and o['modifier']==f'day {day}' and o['data_status']=='author_provided']
    check(f"Enomoto Day {day}: unique extraction row",len(extracted)==1)
    e=extracted[0]
    check(f"Enomoto Day {day}: source statistics preserved",all(float(e[k])==float(row[i]) for k,i in [('n',1),('value',3),('dispersion_value',4),('ci_low',6),('ci_high',7),('p_value',10)]))
    check(f"Enomoto Day {day}: SE = paired SD / sqrt(n)",abs(float(row[4])/math.sqrt(float(row[1]))-float(row[5]))<1e-5)
    check(f"Enomoto Day {day}: t = difference / SE; df = n-1",abs(float(row[3])/float(row[5])-float(row[8]))<0.00051 and float(row[9])==float(row[1])-1)
rob=load_workbook(ROOT/'risk_of_bias.xlsx',data_only=True)
formulas=load_workbook(ROOT/'risk_of_bias.xlsx',data_only=False)
assessment=list(rob['Assessments'].values)[1:]
ranks=['Definitely Low','Probably Low','Probably High','Definitely High']
check("17 studies with one assessment per item",len(assessment)==119 and set(r[0] for r in assessment)==study_ids and all(Counter(r[3] for r in assessment if r[0]==study)==Counter(range(1,8)) for study in study_ids))
check("Valid risk judgments",all(r[6] in ranks for r in assessment))
check("Risk-of-bias evidence and source retained",all(r[7] and r[8] for r in assessment))
for row in range(2,6):
    for item in range(1,8):
        check(f"RoB summary {ranks[row-2]}, item {item}",rob['Summary'].cell(row,item+1).value==sum(r[3]==item and r[6]==ranks[row-2] for r in assessment))
        check(f"RoB summary linked: row {row}, item {item}",formulas['Summary'].cell(row,item+1).data_type=='f')
highest=[]
for row in range(10,27):
    study=rob['Summary'].cell(row,1).value
    expected=max((r[6] for r in assessment if r[0]==study and r[3]<=6),key=ranks.index)
    highest.append(expected)
    check(f"Highest core item: {study}",rob['Summary'].cell(row,2).value==expected and formulas['Summary'].cell(row,2).data_type=='f')
check("Highest-core distribution",[rob['Summary'].cell(r,9).value for r in range(2,6)]==[highest.count(label) for label in ranks])
for file in ['extraction.xlsx','risk_of_bias.xlsx']:
    w=load_workbook(ROOT/file,data_only=False)
    check(f"No broken cell references: {file}",not any(c.data_type=='e' or (c.data_type=='f' and '#REF!' in c.value) for s in w for row in s for c in row))
destination=ROOT/'meta_analysis/outputs/data_validation.csv'
with destination.open('w',encoding='utf-8',newline='') as f:
    writer=csv.writer(f,lineterminator='\n');writer.writerow(['check','status']);writer.writerows(checks)
failures=[name for name,status in checks if status=='FAIL']
if failures:
    raise SystemExit('\n'.join(failures))
print(f'{len(checks)} data and workbook checks passed.')
